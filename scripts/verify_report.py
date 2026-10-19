import docx, os
path = r'd:\BC project\VRS_Mini_Project_Report_Template.docx'
doc = docx.Document(path)
parts = doc.part.related_parts
imgs = [k for k,v in parts.items() if 'image' in getattr(v,'content_type','')]
h2 = [p.text for p in doc.paragraphs if p.style.name=='Heading 2']
h3 = [p.text for p in doc.paragraphs if p.style.name=='Heading 3']
li = [p.text for p in doc.paragraphs if p.style.name=='List Paragraph']
print(f'Size: {os.path.getsize(path)//1024} KB')
print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Tables: {len(doc.tables)}')
print(f'Images (logo): {len(imgs)}')
print()
print('Heading 2 (centered):', h2)
print('Heading 3 (sections):', h3)
print('List Paragraph items:', li)
