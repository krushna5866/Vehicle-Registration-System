import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
sec1 = doc.sections[0]
doc.add_paragraph('Page 1 (Front Matter)')
sec2 = doc.add_section()
sec2.header.is_linked_to_previous = False
sec2.footer.is_linked_to_previous = False
sec2._sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="4"/>'))
pf = sec2.footer.paragraphs[0]
pf.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
pf._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>'))
doc.add_paragraph('Page 4 Content')
doc.save('test_pg.docx')
print('OK')
