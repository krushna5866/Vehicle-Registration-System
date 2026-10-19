"""
build_template_report.py
Rebuilds the VRS Mini Project Report exactly matching the user-provided MGM template.
Font: Times New Roman throughout (matching template)
Logo: extracted from template
"""

import os, copy
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from lxml import etree

LOGO = os.path.abspath(r"d:\BC project\scripts\template_logorId5.png")
OUTPUT = os.path.abspath(r"d:\BC project\VRS_Mini_Project_Report_Template.docx")
FONT = "Times New Roman"

def build():
    doc = Document()

    # ── Page layout (matches template) ─────────────────────────────────
    sec = doc.sections[0]
    sec.page_width  = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin    = Pt(67)   # ~0.93"
    sec.bottom_margin = Pt(14)   # ~0.19"
    sec.left_margin   = Pt(61)   # ~0.85"
    sec.right_margin  = Pt(60)   # ~0.83"

    # ── Helpers ─────────────────────────────────────────────────────────
    def p_add(text="", size=12, bold=False, italic=False,
              align=WD_ALIGN_PARAGRAPH.LEFT,
              color=None, space_before=0, space_after=6,
              underline=False, line_spacing=None):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after  = Pt(space_after)
        if line_spacing:
            p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.underline = underline
            if color:
                r.font.color.rgb = color
        return p

    def center(text, size=12, bold=False, space_before=0, space_after=6, color=None, underline=False):
        return p_add(text, size, bold, align=WD_ALIGN_PARAGRAPH.CENTER,
                     space_before=space_before, space_after=space_after,
                     color=color, underline=underline)

    def justify(text, size=12, bold=False, space_before=0, space_after=6, line_spacing=None):
        return p_add(text, size, bold, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                     space_before=space_before, space_after=space_after,
                     line_spacing=line_spacing)

    def heading2(text):
        p = doc.add_paragraph(style="Heading 2")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after  = Pt(6)
        r = p.add_run(text)
        r.font.name = FONT
        return p

    def heading3(text):
        p = doc.add_paragraph(style="Heading 3")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(text)
        r.font.name = FONT
        return p

    def list_item(text, size=12, bold=True):
        p = doc.add_paragraph(style="List Paragraph")
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after  = Pt(2)
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        return p

    def body_text(text, size=12, bold=False):
        p = doc.add_paragraph(style="Body Text")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(6)
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        return p

    def tbl_border(tbl, color="000000"):
        tbl._tbl.tblPr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="6" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="6" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="4" w:color="{color}"/>'
            f'</w:tblBorders>'))

    def cell_text(cell, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after  = Pt(2)
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold

    def horizontal_line():
        """Add a simple horizontal line paragraph."""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="000000"/>'
            f'</w:pBdr>')
        pPr.append(pBdr)
        return p

    # ====================================================================
    # COVER PAGE
    # ====================================================================

    # Logo
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after  = Pt(6)
    run_logo = p_logo.add_run()
    run_logo.add_picture(LOGO, width=Inches(1.4))

    center("School of Engineering and Technology, MGM University", 16, bold=True, space_after=2)
    center("Chh. Sambhajinagar", 16, bold=True, space_after=4)
    horizontal_line()
    center("Academic Year 2025-2026", 13, bold=True, space_before=4, space_after=8)

    center("Mini Project Report", 12, space_after=2)
    center("On", 12, space_after=8)

    center('"Decentralized Vehicle Registration & Immutable Title Management System Using Blockchain Technology"',
           16, bold=True, space_after=10)

    center("Submitted in partial fulfillment for Sixth Semester of B.Tech (Integrated) in", 12, space_after=2)
    center("COMPUTER SCIENCE AND ENGINEERING", 13, bold=True, space_after=2)
    center("(IoT, Cyber Security Including Block Chain Technology)", 13, bold=True, space_after=10)

    center("By", 12, space_after=6)

    # Student table
    tbl_s = doc.add_table(rows=2, cols=2)
    tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_s)
    tbl_s.rows[0].cells[0].width = Inches(3.5)
    tbl_s.rows[0].cells[1].width = Inches(2.5)
    cell_text(tbl_s.rows[0].cells[0], "Rohit Dhananjay Kale", 12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(tbl_s.rows[0].cells[1], "202356112007", 12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(tbl_s.rows[1].cells[0], "Soham Bhagwan Deshmukh", 12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(tbl_s.rows[1].cells[1], "202356112008", 12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_add()  # spacer
    center("Under the guidance of", 12, space_before=6, space_after=2)
    center("Ms. Bhagat C. B.", 13, bold=True, space_after=2)
    center("(Asst. Professor)", 12, space_after=14)

    horizontal_line()
    center("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", 14, bold=True,
           space_before=8, space_after=2)
    center("School of Engineering and Technology, MGM University, Chh. Sambhajinagar", 12,
           bold=True, space_after=2)

    doc.add_page_break()

    # ====================================================================
    # CERTIFICATE
    # ====================================================================
    # Logo on certificate page
    p_logo2 = doc.add_paragraph()
    p_logo2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo2.paragraph_format.space_before = Pt(0)
    p_logo2.paragraph_format.space_after  = Pt(6)
    p_logo2.add_run().add_picture(LOGO, width=Inches(1.4))

    center("School of Engineering and Technology, MGM University", 16, bold=True, space_after=2)
    center("Chh. Sambhajinagar", 16, bold=True, space_after=4)
    horizontal_line()
    center("CERTIFICATE", 14, bold=True, space_before=12, space_after=12, underline=True)

    # Certificate text
    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert.paragraph_format.space_before = Pt(0)
    p_cert.paragraph_format.space_after  = Pt(10)
    p_cert.paragraph_format.line_spacing = Pt(22)

    def cr(txt, bold=False):
        r = p_cert.add_run(txt)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.font.bold = bold

    cr("This is to certify that ")
    cr("Rohit Dhananjay Kale", True); cr(", ")
    cr("Soham Bhagwan Deshmukh", True)
    cr(" with PRN No. ")
    cr("202356112007", True); cr(", ")
    cr("202356112008", True)
    cr(" have successfully completed their Mini-Project titled ")
    cr('"Decentralized Vehicle Registration & Immutable Title Management System Using Blockchain Technology"', True)
    cr(" in partial fulfillment for the Sixth Semester of B.Tech (Integrated) in Computer Science and Engineering "
       "(IoT, Cyber Security Including Block Chain Technology) from School of Engineering and Technology, "
       "MGM University, Chh. Sambhajinagar during the academic year 2025-2026.")

    p_add(space_before=18)  # spacer

    # Signature table
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sig.paragraph_format.space_before = Pt(30)
    def sr(txt, bold=False):
        r = p_sig.add_run(txt)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.font.bold = bold
    sr("Ms. Bhagat C. B.", True)
    sr("\t\t\t\t")
    sr("Ms. Rathi S. R.", True)

    p_role = doc.add_paragraph()
    p_role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    def rr(txt, bold=True):
        r = p_role.add_run(txt)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.font.bold = bold
    rr("Guide")
    rr("\t\t\t\t\t\t")
    rr("HOD")

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_dept = p_dept.add_run("         CSE Department\t\t\t         CSE Department")
    r_dept.font.name = FONT
    r_dept.font.size = Pt(12)

    doc.add_page_break()

    # ====================================================================
    # CONTENT (Table of Contents)
    # ====================================================================
    heading2("CONTENT")

    toc_data = [
        ("1",  "Abstract",                                    "01"),
        ("2",  "Introduction",                                "02"),
        ("3",  "Course Outcomes Integrated",                  "03"),
        ("4",  "Action Plan",                                 "04"),
        ("5",  "Literature Review",                           "05"),
        ("6",  "Proposed Methodology",                        "06"),
        ("7",  "Implementation",                              "07"),
        ("8",  "Outputs of the Mini-Project",                 "09"),
        ("9",  "Skill Developed / Learning out of this Mini-Project", "11"),
        ("10", "Applications of this Mini-Project",           "12"),
        ("11", "Area of Future Improvement",                  "13"),
        ("12", "Conclusion",                                  "14"),
    ]
    tbl_toc = doc.add_table(rows=len(toc_data)+1, cols=3)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_toc)
    tw = [Inches(0.8), Inches(4.7), Inches(0.9)]

    for ci, hd in enumerate(["Sr. No", "Title", "Page. No."]):
        cell_text(tbl_toc.rows[0].cells[ci], hd, 12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    for i, (n, t, pg) in enumerate(toc_data):
        r = tbl_toc.rows[i+1]
        cell_text(r.cells[0], n, 12, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(r.cells[1], t, 12, align=WD_ALIGN_PARAGRAPH.LEFT)
        cell_text(r.cells[2], pg, 12, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()

    # ====================================================================
    # 1. ABSTRACT
    # ====================================================================
    list_item("Abstract")

    body_text(
        "Vehicle registration in India is handled by Regional Transport Offices (RTOs) using centralized "
        "government databases. These systems face serious problems — a corrupt official can secretly edit records, "
        "a server failure can bring everything down, and buyers of second-hand cars have no independent way to "
        "verify a vehicle's real ownership history. Fraudsters exploit these gaps by selling stolen or salvage "
        "vehicles using forged Registration Certificates."
    )
    body_text(
        "This mini-project builds a Decentralized Vehicle Registration System on the Ethereum blockchain. "
        "Every vehicle title is stored as a digital token (ERC-721 NFT) that no one can secretly alter. "
        "To protect citizen privacy, only Keccak-256 fingerprint hashes of the chassis number (VIN) and "
        "number plate are stored on-chain — no names, phone numbers, or addresses."
    )
    body_text(
        "The system enforces strict role separation: an RTO Registrar registers vehicles and approves "
        "ownership transfers through the official two-step process, while an RTO Enforcement Officer "
        "handles vehicle scrappage. Neither can interfere with the other's duties, and no authority "
        "wallet can ever own a vehicle title. A total of 40 automated tests were run, all passing "
        "successfully, including 16 negative security attack scenarios."
    )
    doc.add_page_break()

    # ====================================================================
    # 2. INTRODUCTION
    # ====================================================================
    list_item("Introduction")

    body_text(
        "In India, every vehicle on the road must be registered with the local Regional Transport Office (RTO). "
        "The Registration Certificate (RC) is the official proof of ownership. If you want to sell a car, "
        "buy a second-hand car, or get vehicle insurance, the RC is the most important document involved."
    )
    body_text(
        "Today, RTOs maintain these records on centralized computer systems, such as the national Vahan portal. "
        "While this is better than old paper registers, it still has serious weaknesses that affect ordinary citizens."
    )
    body_text("Problems with the Current System:", bold=True)
    body_text(
        "1. Fake Vehicles and Forged Documents: Criminals steal a car, duplicate its chassis plate, and obtain "
        "fake RCs. Innocent buyers purchase these vehicles and later lose them to law enforcement with no refund.\n"
        "2. Database Tampering: A corrupt official or hacker can change an owner's name, remove a bank loan "
        "record, or alter a vehicle's status. There is no permanent audit trail.\n"
        "3. No Transparent History: When buying a used car, there is no independent public way to verify "
        "the vehicle's full ownership history. Buyers must trust paperwork from the seller.\n"
        "4. Salvage Car Fraud: When a car is totaled in an accident, the chassis is sometimes secretly rebuilt "
        "and resold. Buyers have no way to know this happened.\n"
        "5. Authority Abuse: Government staff can use their admin access to manipulate records. There are no "
        "automated technical checks to prevent this."
    )
    body_text("How This Project Solves These Problems:", bold=True)
    body_text(
        "Our system stores vehicle titles as blockchain tokens (ERC-721 NFTs) — no record can ever be changed "
        "or deleted once written. Only Keccak-256 fingerprints of VINs and plates are stored publicly; citizen "
        "data stays in private encrypted storage. Every ownership change is permanently timestamped on-chain. "
        "RTO Registrars can only register and transfer; RTO Officers can only scrap. Neither role can do the "
        "other's job, and no RTO wallet can ever own a vehicle. Once a vehicle is scrapped, it is permanently "
        "locked — enforced by code, not by human honesty."
    )
    doc.add_page_break()

    # ====================================================================
    # 3. COURSE OUTCOME INTEGRATED
    # ====================================================================
    list_item("Course Outcome Integrated")

    p_co_label = doc.add_paragraph()
    p_co_label.paragraph_format.space_before = Pt(4)
    p_co_label.paragraph_format.space_after  = Pt(6)
    r_co = p_co_label.add_run("          List of Course Outcomes of Course which are Related to Mini Project")
    r_co.font.name = FONT
    r_co.font.size = Pt(12)
    r_co.font.bold = True

    co_data = [
        ("CO1", "Cryptographic Foundations",
         "Used Keccak-256 hashing to protect vehicle identity data. Digital signatures are used for wallet login without exposing personal details."),
        ("CO2", "Smart Contract Development",
         "Wrote a complete, deployable Solidity 0.8.24 smart contract using OpenZeppelin libraries for ERC-721 tokens, access control, and emergency pause."),
        ("CO3", "Role-Based Access Control",
         "Designed separate roles — Admin, RTO Registrar, RTO Officer, and Citizen — each with strict permissions enforced at the smart contract level."),
        ("CO4", "DApp Development",
         "Built a full React 18 + Vite web application that connects to Ethereum through Ethers.js v6, with route guards and live transaction handling."),
        ("CO5", "Security Testing",
         "Wrote 40 automated Hardhat test cases, covering 16 negative security invariants — unauthorized access, privilege escalation, and state machine violations."),
    ]
    tbl_co = doc.add_table(rows=len(co_data)+1, cols=3)
    tbl_co.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_co)
    cw = [Inches(0.8), Inches(2.0), Inches(3.7)]
    for ci, hd in enumerate(["CO", "Topic", "How This Project Uses It"]):
        cell_text(tbl_co.rows[0].cells[ci], hd, 12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, (co, topic, how) in enumerate(co_data):
        r = tbl_co.rows[i+1]
        cell_text(r.cells[0], co, 12, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(r.cells[1], topic, 12, align=WD_ALIGN_PARAGRAPH.LEFT)
        cell_text(r.cells[2], how, 12, align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_page_break()

    # ====================================================================
    # 4. ACTION PLAN
    # ====================================================================
    list_item("Action Plan")

    body_text(
        "The project was completed in six phases from August to October 2025. Both team members divided "
        "responsibilities based on their strengths and worked together for testing and documentation."
    )

    act_data = [
        ("1", "Study of RTO workflows (Form 29/30) and System Architecture Design",
         "01-Aug-2025", "12-Aug-2025", "Both Members"),
        ("2", "Smart Contract Development: VehicleRegistry.sol (ERC-721 and Role-Based Access Control)",
         "13-Aug-2025", "28-Aug-2025", "Rohit Kale"),
        ("3", "Privacy Layer: Keccak-256 Commitment Scheme and Off-Chain Encrypted Storage",
         "29-Aug-2025", "10-Sep-2025", "Soham Deshmukh"),
        ("4", "Security Boundaries: Two-Step Transfer Workflow and Scrappage State Machine",
         "11-Sep-2025", "24-Sep-2025", "Rohit Kale"),
        ("5", "Frontend Development: React/Vite Web Application, Citizen and RTO Portals",
         "25-Sep-2025", "12-Oct-2025", "Soham Deshmukh"),
        ("6", "Testing (40 Test Cases), Live Blockchain Demonstration and Report Writing",
         "13-Oct-2025", "28-Oct-2025", "Both Members"),
    ]
    tbl_a = doc.add_table(rows=len(act_data)+1, cols=5)
    tbl_a.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_a)
    aw = [Inches(0.5), Inches(2.9), Inches(1.0), Inches(1.0), Inches(1.1)]
    for ci, hd in enumerate(["S. No.", "Details of Activity", "Planned Start Date", "Planned Finish Date", "Name of Responsible Team Members"]):
        cell_text(tbl_a.rows[0].cells[ci], hd, 11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, (n, act, sd, ed, who) in enumerate(act_data):
        rw = tbl_a.rows[i+1]
        cell_text(rw.cells[0], n, 11, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(rw.cells[1], act, 11, align=WD_ALIGN_PARAGRAPH.LEFT)
        cell_text(rw.cells[2], sd, 11, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(rw.cells[3], ed, 11, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(rw.cells[4], who, 11, align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_page_break()

    # ====================================================================
    # 5. LITERATURE REVIEW
    # ====================================================================
    heading3("Literature Review")

    body_text(
        "We studied existing vehicle registration systems and earlier blockchain-based proposals before "
        "designing our solution."
    )
    body_text(
        "The Government of India's Vahan portal is the national vehicle registration database used by all RTOs. "
        "It improved record-keeping compared to manual registers and allowed online tax payment. However, since "
        "all data lives on central government servers, it remains vulnerable to cyber attacks, server failures, "
        "and insider misuse. When a record is altered, there is no automatic way to detect or reverse it."
    )
    body_text(
        "Several university projects have attempted to use blockchain for vehicle registration. Most of them "
        "had two major flaws. First, they stored citizens' real names, phone numbers, and addresses directly "
        "on the public blockchain — a serious privacy violation since anyone can read blockchain data. Second, "
        "they allowed citizens to transfer vehicle ownership directly, bypassing all RTO checks — which is "
        "illegal under Indian motor vehicle law."
    )

    # Comparison table
    body_text("Comparison of Existing Systems vs Our Approach:", bold=True)
    comp = [
        ("Tamper-Proof Records",        "No — can be edited by admins",  "Partially — depends on design",  "Yes — blockchain is immutable"),
        ("Privacy Protection",           "Moderate — central server risk", "No — PII exposed on-chain",      "Yes — only hashes stored on-chain"),
        ("Transparent Ownership History","No — not publicly accessible",   "Partial — basic records only",   "Yes — complete history on blockchain"),
        ("Legal Transfer Compliance",    "Manual paperwork, slow",          "Not enforced",                   "Yes — RTO Registrar must approve"),
        ("Role Separation",              "Loose — single admin",            "None — single admin",            "Strict — Registrar vs Officer vs Admin"),
        ("Scrappage Enforcement",        "Paper certificate, not enforced", "Boolean flag — easily bypassed", "Permanent state lock — cannot be undone"),
    ]
    tbl_c = doc.add_table(rows=len(comp)+1, cols=4)
    tbl_c.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_c)
    kcw = [Inches(1.5), Inches(1.5), Inches(1.5), Inches(2.0)]
    for ci, hd in enumerate(["Feature", "Traditional Database", "Earlier Blockchain Projects", "Our System"]):
        cell_text(tbl_c.rows[0].cells[ci], hd, 11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, row_data in enumerate(comp):
        rw = tbl_c.rows[i+1]
        for ci, val in enumerate(row_data):
            cell_text(rw.cells[ci], val, 11, align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_page_break()

    # ====================================================================
    # 6. PROPOSED METHODOLOGY
    # ====================================================================
    heading3("Proposed Methodology")

    body_text(
        "Our system is built in four layers. Users interact at the frontend web application, which enforces "
        "security checks before any data reaches the blockchain. Sensitive vehicle identifiers are converted "
        "into hash fingerprints for privacy, and the final records are permanently stored in the smart contract."
    )
    body_text("The Four Layers of the System:", bold=True)
    body_text(
        "Layer 1 — Users: Ordinary citizens, used-car buyers, RTO Registrars, and RTO Officers. "
        "Each type of user has their own dedicated portal in the web application.\n\n"
        "Layer 2 — Web Application and Security: A React 18 web application built with Vite. "
        "Each portal has a security guard (CitizenGuard or RtoGuard) that blocks the wrong user type from entering. "
        "If you switch your wallet in MetaMask, you are immediately logged out.\n\n"
        "Layer 3 — Privacy Conversion: Before any data is sent to the blockchain, sensitive vehicle identifiers "
        "(VIN and plate number) are converted into Keccak-256 fingerprint hashes. Names, addresses, and phone "
        "numbers stay in private encrypted storage and never go on the public blockchain.\n\n"
        "Layer 4 — Blockchain (Smart Contract): The Ethereum smart contract (VehicleRegistry.sol) where vehicle "
        "titles are created, ownership is transferred, vehicles are scrapped, and history is permanently recorded."
    )
    body_text("How Hashing Protects Privacy:", bold=True)
    body_text(
        "A Keccak-256 hash works like a fingerprint scanner. The chassis number (e.g., MA3ERLF3S00123456) "
        "is hashed to a unique 32-byte code stored on-chain. The plate number (e.g., MH20DE1234) is also "
        "hashed separately. A police officer can verify a vehicle is genuine by entering the plate — "
        "the system hashes it and checks against the stored fingerprint. But no one can reverse the hash "
        "to find out who owns the car."
    )
    body_text("Two RTO Roles — Why Keep Them Separate?:", bold=True)
    body_text(
        "In a traditional system, one senior admin often has all powers — registering cars, transferring them, "
        "and scrapping them. This is dangerous because a single corrupt official can misuse everything. "
        "We split the work into two clearly different roles:\n\n"
        "RTO Registrar: Can register new vehicles and approve ownership transfers. Cannot scrap vehicles. "
        "Cannot own vehicle titles.\n\n"
        "RTO Enforcement Officer: Can mark vehicles as Scrapped. Cannot register vehicles. Cannot transfer titles. "
        "Cannot own vehicle titles.\n\n"
        "Anti-Corruption Rule: Neither the Admin nor any RTO wallet can ever receive a vehicle title. "
        "If anyone tries to transfer a car to an RTO wallet, the smart contract automatically rejects it."
    )
    body_text("How Ownership Transfer Works (Step by Step):", bold=True)
    body_text(
        "Step 1 — Citizen Initiates: The current vehicle owner logs into the Citizen Portal and submits a "
        "transfer request with the buyer's wallet address. The car stays in the seller's name at this point.\n\n"
        "Step 2 — RTO Registrar Approves: The RTO Registrar reviews the request, verifies that all legal "
        "conditions are met (taxes paid, no fines pending, loan cleared), and calls the transfer function on "
        "the smart contract.\n\n"
        "Step 3 — Title Moves, History Updated: The smart contract moves the digital title token to the buyer's "
        "wallet and permanently records the buyer in the vehicle's ownership history with a timestamp."
    )
    body_text("Vehicle Scrappage — Permanent Terminal State:", bold=True)
    body_text(
        "When a vehicle is totaled in a crash or reaches the end of its life, an RTO Enforcement Officer can "
        "mark it as Scrapped. This is irreversible on the blockchain. All pending transfer requests are "
        "immediately cancelled. The vehicle title is permanently locked — no transfers, no re-registrations, ever. "
        "This prevents the common fraud of rebuilding crashed cars and selling them to unknowing buyers."
    )
    doc.add_page_break()

    # ====================================================================
    # 7. IMPLEMENTATION
    # ====================================================================
    heading3("Implementation")

    body_text("Smart Contract — VehicleRegistry.sol:", bold=True)
    body_text(
        "Written in Solidity 0.8.24 and deployed on a local Ethereum node (Hardhat). "
        "Built using OpenZeppelin v5.0 trusted library contracts: ERC721 (vehicle title tokens), "
        "AccessControl (role management), and Pausable (emergency shutdown)."
    )
    body_text("Main data stored for each vehicle on the blockchain:", bold=False)
    body_text(
        "- vehicleId: Unique ID derived from the hash of VIN + Plate\n"
        "- vinCommitment: Keccak-256 hash of chassis number\n"
        "- registrationCommitment: Keccak-256 hash of number plate\n"
        "- state: Active (0) or Scrapped (1)\n"
        "- titleHolder: Current owner's wallet address\n"
        "- metadataHash: Hash of vehicle specs (make, model, year)\n"
        "- registeredAt: Timestamp of first registration"
    )
    body_text("Key Smart Contract Functions:", bold=True)

    tbl_f = doc.add_table(rows=5, cols=3)
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_f)
    fw = [Inches(1.7), Inches(1.4), Inches(3.4)]
    for ci, hd in enumerate(["Function", "Who Can Call It", "What It Does"]):
        cell_text(tbl_f.rows[0].cells[ci], hd, 11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    funcs = [
        ("registerVehicle(...)",          "RTO Registrar Only", "Creates a new vehicle record and mints an ERC-721 title token for the citizen owner."),
        ("transferOwnershipByRTO(...)",   "RTO Registrar Only", "Moves the title token to the new owner after RTO approval. Records the transfer in ownership history."),
        ("initiateTransfer(...)",         "Citizen (Title Holder)", "Submits a transfer request. RTO Registrar must still approve for the title to actually move."),
        ("updateVehicleState(...)",       "RTO Officer Only",   "Marks the vehicle as Scrapped permanently. Locks the title forever. This action cannot be undone."),
    ]
    for i, (fn, who, what) in enumerate(funcs):
        rw = tbl_f.rows[i+1]
        cell_text(rw.cells[0], fn, 11, align=WD_ALIGN_PARAGRAPH.LEFT)
        cell_text(rw.cells[1], who, 11, align=WD_ALIGN_PARAGRAPH.LEFT)
        cell_text(rw.cells[2], what, 11, align=WD_ALIGN_PARAGRAPH.LEFT)

    body_text("Security Hook — No Privilege Owner Rule:", bold=True)
    body_text(
        "We override OpenZeppelin's internal token transfer function. Every time any token transfer is attempted "
        "— even using Ethereum's transferFrom directly — it first checks whether the recipient is a privileged "
        "address (RTO or Admin wallet). If so, the transaction is automatically rejected with the error: "
        "PrivilegedAddressCannotOwnVehicle. This rule cannot be bypassed by anyone, even via direct contract calls."
    )
    body_text("Web Application:", bold=True)
    body_text(
        "The frontend is a React 18 + Vite application using Ethers.js v6 to communicate with the Ethereum network. "
        "It has four portals: a Public Search portal (anyone can verify a vehicle), a Citizen Dashboard "
        "(owners view their title as an HSRP plate display and request transfers), an RTO Registrar Portal "
        "(register vehicles and process transfers), and an RTO Enforcement Officer Portal (view details and "
        "permanently scrap vehicles)."
    )
    doc.add_page_break()

    # ====================================================================
    # 8. OUTPUT OF MINI-PROJECT
    # ====================================================================
    list_item("Outputs of the Mini-Project")

    body_text(
        "To prove the system works correctly and securely, we wrote 40 automated test cases using Hardhat "
        "and Chai. All 40 tests passed on the final run. The following groups were covered:"
    )

    tbl_t = doc.add_table(rows=9, cols=3)
    tbl_t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_t)
    ttw = [Inches(0.5), Inches(3.5), Inches(2.5)]
    for ci, hd in enumerate(["#", "Test Group", "What It Verifies"]):
        cell_text(tbl_t.rows[0].cells[ci], hd, 11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    tgroups = [
        ("1", "Role-Based Access Control & RTO Role Boundaries", "Admin, Registrar, and Officer each have only what they need — nothing more"),
        ("2", "On-Chain Privacy (No Duplicate VIN / Plate)",      "Duplicate registrations are blocked at the hash level"),
        ("3", "Two-Step Transfer Workflow",                        "Citizen can request; only RTO Registrar can approve and execute the transfer"),
        ("4", "Vehicle Scrappage & Terminal State",                "Scrapped vehicles cannot be transferred or re-registered ever again"),
        ("5", "Ownership History Tracking",                        "All past owners are permanently recorded with timestamps on-chain"),
        ("6", "Emergency Pause Controls",                          "Admin can pause the system; no state changes work while paused"),
        ("7", "Event Privacy",                                     "Blockchain events contain only hashes — no raw personal data exposed"),
        ("8", "16 Negative Security Invariants",                   "16 different attack scenarios are all correctly blocked and rejected"),
    ]
    for i, (n, grp, what) in enumerate(tgroups):
        rw = tbl_t.rows[i+1]
        cell_text(rw.cells[0], n, 11, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(rw.cells[1], grp, 11, align=WD_ALIGN_PARAGRAPH.LEFT)
        cell_text(rw.cells[2], what, 11, align=WD_ALIGN_PARAGRAPH.LEFT)

    body_text("Test Result Summary:", bold=True)
    body_text(
        "Total Test Cases: 40   |   Passed: 40   |   Failed: 0   |   Time: 5 seconds\n"
        "All 16 negative security invariants: BLOCKED\n"
        "Ownership history tracking: VERIFIED\n"
        "Scrappage terminal state: ENFORCED\n"
        "Emergency pause: WORKING\n"
        "Role boundaries (Registrar / Officer): ENFORCED"
    )
    body_text("Live Blockchain Test:", bold=True)
    body_text(
        "We ran a live demonstration on a local Ethereum node (Chain ID: 31337). "
        "The contract was deployed and real transactions were executed to prove end-to-end functionality:\n"
        "- Citizen wallet blocked from registering a vehicle\n"
        "- Duplicate plate commitment rejected on-chain\n"
        "- Ownership officially transferred via RTO Registrar\n"
        "- RTO Officer blocked from calling transfer functions\n"
        "- Full ownership history verified on-chain (3 records)\n"
        "- Vehicle scrapped by RTO Officer — state locked as SCRAPPED\n"
        "- Transfer of scrapped vehicle blocked with error\n"
        "- RTO wallet blocked from receiving a vehicle title\n"
        "- Emergency pause activated — all operations blocked\n"
        "All proofs verified successfully on live chain."
    )
    doc.add_page_break()

    # ====================================================================
    # 9. SKILLS DEVELOPED
    # ====================================================================
    list_item(" Skill Developed / Learning out of this Mini-Project")

    body_text(
        "          Working on this project gave us hands-on experience with technologies that go well beyond "
        "standard textbook exercises. We learned to write real, deployable smart contracts in Solidity and "
        "understood how Ethereum enforces rules automatically without any middleman. We practically applied "
        "Keccak-256 hashing to solve the real privacy problem of public blockchains, and now understand the "
        "difference between storing data publicly and proving it securely. We designed and enforced role-based "
        "permissions at the code level, understanding the Principle of Least Privilege in a government system context. "
        "We built a complete working Web3 application that connects to a blockchain in real time, handles wallet "
        "authentication, manages user sessions, and responds to live blockchain events. Most importantly, we "
        "learned test-driven security — writing 40 test cases and actively thinking like an attacker to find "
        "and fix vulnerabilities before they become real problems."
    )
    doc.add_page_break()

    # ====================================================================
    # 10. APPLICATIONS
    # ====================================================================
    heading3("Applications of this Mini-Project")

    body_text(
        "The architecture we built can be applied to solve real problems in India and worldwide. State RTOs "
        "can use this system to eliminate forged Registration Certificates, duplicate registrations, and "
        "insider tampering across state borders. Used vehicle marketplaces like Cars24 and Spinny can "
        "instantly verify a vehicle's complete ownership history and scrappage status without manual RTO "
        "paperwork. Insurance companies can trigger on-chain scrappage when declaring a total loss, "
        "preventing the salvage car from ever being resold to an unknowing buyer. Banks giving car loans "
        "can place a digital lien on the vehicle title during the loan period — the title cannot be "
        "transferred until the loan is fully repaid. With minor modifications, the same system can also "
        "be applied to land record management, machinery certification, and any domain where tamper-proof "
        "ownership history is critical."
    )
    doc.add_page_break()

    # ====================================================================
    # 11. FUTURE IMPROVEMENT
    # ====================================================================
    heading3("Area of Future Improvement")

    body_text(
        "Our current system is a solid, working prototype. Several enhancements can make it production-ready "
        "for nationwide government use."
    )
    body_text(
        "Zero-Knowledge Proof Integration: A citizen could prove their vehicle has valid registration to a "
        "police officer without revealing their wallet address or identity, using zk-SNARK cryptography. "
        "This would provide even stronger privacy guarantees."
    )
    body_text(
        "IPFS Storage for Documents: Vehicle inspection photos, fitness certificates, and insurance documents "
        "can be stored on IPFS (a distributed file network) instead of a central server, making the entire "
        "system fully decentralized."
    )
    body_text(
        "IoT and FASTag Integration: Vehicles equipped with OBD-II chips and FASTag transponders could directly "
        "submit signed, tamper-proof odometer readings to the blockchain, eliminating odometer rollback fraud "
        "which is very common in the used-car market."
    )
    body_text(
        "Layer-2 Deployment for Scale: Deploying on Ethereum Layer-2 networks like Polygon or Arbitrum would "
        "reduce transaction costs from a few dollars to a fraction of a rupee, making it practical and "
        "affordable for nationwide adoption."
    )
    doc.add_page_break()

    # ====================================================================
    # 12. CONCLUSION
    # ====================================================================
    p_conc_hd = doc.add_paragraph(style="Body Text")
    p_conc_hd.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_conc_hd = p_conc_hd.add_run("Conclusion")
    r_conc_hd.font.name = FONT
    r_conc_hd.font.size = Pt(12)
    r_conc_hd.font.bold = True

    body_text(
        "This mini-project shows that blockchain technology can solve real, everyday problems that affect "
        "ordinary citizens — not just cryptocurrency transactions. By building a vehicle registration system "
        "on Ethereum, we created records that nobody can secretly alter, a complete ownership history that "
        "anyone can independently verify, and a scrappage mechanism that permanently removes dangerous salvage "
        "cars from circulation — all while keeping citizens' personal details fully private through "
        "cryptographic hashing."
    )
    body_text(
        "We also demonstrated that government administrative systems can be made tamper-proof by design, "
        "using smart contract rules instead of relying entirely on human honesty. The strict separation "
        "between the RTO Registrar and the RTO Enforcement Officer ensures that no single official has "
        "unchecked power to manipulate vehicle records. Since neither role can own a vehicle, there is no "
        "way for an authority to misuse their position to transfer a vehicle to themselves."
    )
    body_text(
        "Our 40 automated tests, all passing at 100%, and the live on-chain demonstration prove that this "
        "is not just a theoretical concept — it is a working, deployable system that addresses real gaps "
        "in India's current vehicle registration infrastructure."
    )

    # ── Save ────────────────────────────────────────────────────────────
    doc.save(OUTPUT)
    size = os.path.getsize(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Size:  {size // 1024} KB")

if __name__ == "__main__":
    build()
