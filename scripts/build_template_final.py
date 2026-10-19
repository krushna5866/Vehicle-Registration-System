"""
build_template_final.py
Refined script to modify user's template (media_1790946874497.docx) with:
- Exact starting pages (Cover, Certificate, CONTENT) from template
- Title updated to Vehicle Registration & Immutable Title Management System
- Certificate updated for Rohit Kale & Soham Deshmukh
- Action Plan table filled with 6 realistic activities
- All body sections replaced with short, very simple, student-friendly content
- Explicit Times New Roman font and clean formatting on all runs
"""

import os
import shutil
import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_PATH = r"C:\Users\Krushna\.gemini\antigravity\brain\05af4329-1e0d-4b3d-af5b-060ca7138a6d\.user_uploaded\media_1790946874497.docx"
OUTPUT_PROJECT = r"d:\BC project\Vehicle_Registration_System_Mini_Project_Report_Template_Final.docx"
HOME = os.path.expanduser("~")
OUTPUT_DESKTOP = os.path.join(HOME, "Desktop", "Vehicle_Registration_System_Mini_Project_Report_Template_Final.docx")
OUTPUT_DOWNLOADS = os.path.join(HOME, "Downloads", "Vehicle_Registration_System_Mini_Project_Report_Template_Final.docx")
OUTPUT_ARTIFACT = r"C:\Users\Krushna\.gemini\antigravity\brain\05af4329-1e0d-4b3d-af5b-060ca7138a6d\Vehicle_Registration_System_Mini_Project_Report_Template_Final.docx"

FONT_NAME = "Times New Roman"

def set_para_text(p, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p.text = ""
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(size)
    r.font.bold = bold
    return r

def build():
    doc = Document(TEMPLATE_PATH)

    # ── 1. Update Cover Page Title (Paragraph 14) ──────────────────────
    p14 = doc.paragraphs[14]
    p14.text = ""
    p14.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r14 = p14.add_run('“Decentralized Vehicle Registration & Immutable Title Management System Using Blockchain Technology”')
    r14.font.name = FONT_NAME
    r14.font.size = Pt(15)
    r14.font.bold = True

    # ── 2. Remove (Data Science) on Cover Page (Paragraph 18) ──────────
    doc.paragraphs[18].text = ""

    # ── 3. Update Certificate Text (Paragraph 50) ──────────────────────
    p50 = doc.paragraphs[50]
    p50.text = ""
    p50.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p50.paragraph_format.line_spacing = 1.25

    def c_run(text, bold=False):
        r = p50.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(11.5)
        r.font.bold = bold

    c_run("This is to certify that ")
    c_run("Rohit Dhananjay Kale", True)
    c_run(", ")
    c_run("Soham Bhagwan Deshmukh", True)
    c_run(" with PRN No. ")
    c_run("202356112007", True)
    c_run(", ")
    c_run("202356112008", True)
    c_run(" has successfully completed their Mini-Project entitled ")
    c_run('“Decentralized Vehicle Registration & Immutable Title Management System Using Blockchain Technology”', True)
    c_run(" in the Course of ")
    c_run('“Blockchain” (Course Code: CSB24PCC452)', True)
    c_run(" in the Sixth Semester during tenure of completing the B. Tech(I) program in Computer Science and Engineering ")
    c_run("(IoT, Cyber Security Including Blockchain Technology)", True)
    c_run(" from School of Engineering and Technology, MGM University institute.")

    # ── 4. Fill Table 2: Action Plan ───────────────────────────────────
    tbl2 = doc.tables[2]
    plan_data = [
        ("1", "Study of RTO workflows (Form 29/30) and System Architecture Design", "18/01/2026", "05/02/2026", "Project Team"),
        ("2", "Smart Contract Development: VehicleRegistry.sol (ERC-721 and RBAC)", "05/02/2026", "18/02/2026", "Rohit Kale"),
        ("3", "Privacy Layer: Keccak-256 Commitments and Off-Chain Storage Vault", "18/02/2026", "28/02/2026", "Soham Deshmukh"),
        ("4", "Security Boundaries: Two-Step Transfer Workflow and Scrappage Lock", "28/02/2026", "08/03/2026", "Rohit Kale"),
        ("5", "Frontend Development: React/Vite Web App, Citizen and RTO Portals", "08/03/2026", "18/03/2026", "Soham Deshmukh"),
        ("6", "Final testing (40 Automated Tests), Invariant Security and Documentation", "18/03/2026", "28/03/2026", "Both Members"),
    ]
    for row_idx, rdata in enumerate(plan_data):
        row = tbl2.rows[row_idx + 1]
        for col_idx, text in enumerate(rdata):
            cell = row.cells[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name = FONT_NAME
            r.font.size = Pt(9.5)

    # ── 5. Replace Body Paragraphs with Short, Very Simple Content ──────

    # Abstract (Paragraph 81)
    set_para_text(
        doc.paragraphs[81],
        "In India, vehicle registration is managed by Regional Transport Offices (RTOs) using centralized databases. "
        "These systems have serious problems like fake registration certificates, stolen cars being sold with duplicate plates, "
        "and records being secretly altered. This mini-project builds a Decentralized Vehicle Registration System using the Ethereum "
        "blockchain. In this system, each vehicle title is stored as a digital NFT that cannot be changed or deleted by anyone. "
        "To protect citizen privacy, only cryptographic hashes (digital fingerprints) of the chassis number and number plate are stored on "
        "the public blockchain, while personal details stay private in an encrypted vault. The system also separates government duties "
        "so that an RTO Registrar registers vehicles and approves transfers, while an RTO Officer only handles vehicle scrappage. "
        "All 40 automated security tests passed successfully with 100% accuracy.",
        size=11.5
    )

    # Introduction (Paragraphs 84, 85, 86)
    set_para_text(
        doc.paragraphs[84],
        "Every vehicle on the road must be registered with the local Regional Transport Office (RTO). "
        "The Registration Certificate (RC) is the official proof of vehicle ownership. It is required to sell a car, "
        "buy insurance, or prove ownership to the police. Today, RTOs store these records in central computer databases like Vahan. "
        "While this is better than old paper records, central databases have big weaknesses that cause real problems for ordinary citizens.",
        size=11.5
    )
    set_para_text(
        doc.paragraphs[85],
        "There are four major problems in the current system: (1) Criminals create fake RC papers for stolen cars and sell them to "
        "innocent buyers; (2) Corrupt insiders or hackers can edit central database records to remove bank loans or change owner names; "
        "(3) Used-car buyers have no easy way to verify the complete, true history of a second-hand car; and (4) Totaled cars from bad "
        "accidents are sometimes secretly rebuilt and sold without the buyer knowing.",
        size=11.5
    )
    set_para_text(
        doc.paragraphs[86],
        "Our project solves these problems using blockchain technology. Digital titles are stored as NFTs that cannot be altered. "
        "Vehicle numbers are protected using cryptographic hashing so private data is safe. A two-step transfer process ensures that "
        "transfers only happen after official RTO approval. Finally, once an officer marks a car as scrapped, it is permanently locked "
        "so salvage cars can never be re-registered or sold again.",
        size=11.5
    )
    for idx in range(87, 90):
        doc.paragraphs[idx].text = ""

    # Course Outcome Integrated (Paragraph 93)
    set_para_text(
        doc.paragraphs[93],
        "Course Code: CSB24PCC452 | Course Name: Blockchain (Credits: 2)\n\n"
        "• CO1: Understand the basic structure, history, and types of blockchain systems and decentralization methods.\n"
        "  - Project Mapping: Replaced centralized client-server RTO databases (Vahan) with an Ethereum decentralized blockchain ledger, eliminating single points of failure and unauthorized database tampering.\n\n"
        "• CO2: Analyze blockchain architecture, smart contracts, and cryptocurrencies in various real-world use cases.\n"
        "  - Project Mapping: Engineered the VehicleRegistry.sol smart contract in Solidity 0.8.24 utilizing OpenZeppelin ERC-721 for vehicle title NFTs, AccessControl for RTO governance, and Pausable for safety.\n\n"
        "• CO3: Apply cryptographic principles and consensus algorithms in Bitcoin transactions and wallet mechanisms.\n"
        "  - Project Mapping: Implemented Keccak-256 cryptographic hashing to store vehicle commitments on-chain while keeping citizen PII off-chain. Implemented Web3 wallet authentication via MetaMask challenge-nonce signing.\n\n"
        "• CO4: Evaluate and design distributed applications using modern blockchain platforms.\n"
        "  - Project Mapping: Designed an end-to-end DApp on Ethereum with a React 18/Vite frontend, Ethers.js v6 integration, two-step title transfers, permanent scrappage locks, and 40 automated tests passing at 100%.",
        size=10.5
    )
    for idx in range(94, 96):
        doc.paragraphs[idx].text = ""

    # Literature Review (Paragraphs 104, 107)
    set_para_text(
        doc.paragraphs[104],
        "We studied existing vehicle registration methods before building our system. In the past, RTOs used paper registers "
        "and physical RC booklets. These were easily damaged, stolen, or forged, and there was no way for different states to talk to each "
        "other. The national Vahan database improved this by putting records online. However, because Vahan is a centralized system, "
        "it has a single point of failure. If the central database is hacked or manipulated from inside, unauthorized changes cannot "
        "be easily detected.",
        size=11.5
    )
    set_para_text(
        doc.paragraphs[107],
        "Earlier student blockchain projects also tried to fix this, but they made two big mistakes. First, they posted citizens' real "
        "names and phone numbers directly onto the public blockchain, which is a major privacy violation. Second, they allowed people to "
        "transfer car ownership directly to each other like sending an email, completely bypassing RTO legal checks. Our project fixes both "
        "flaws: we store only cryptographic hash fingerprints on the blockchain to protect privacy, and we require the RTO Registrar to "
        "verify and approve all ownership transfers.",
        size=11.5
    )

    # Proposed Methodology (Paragraphs 113, 115, 117, 119)
    set_para_text(
        doc.paragraphs[113],
        "Our system is designed in four simple layers: (1) Users — Citizens, buyers, and RTO officers access the system through different "
        "web portals; (2) Frontend & Security — A React web application with built-in route guards that verify user roles before granting access; "
        "(3) Privacy Conversion Layer — Converts chassis numbers (VIN) and number plates into Keccak-256 hash fingerprints before sending them "
        "to the blockchain, keeping citizen personal data in private encrypted storage; and (4) Ethereum Smart Contract (VehicleRegistry.sol) — "
        "Where digital titles are minted, transfers are approved, and ownership history is permanently stored.",
        size=11.5
    )
    set_para_text(
        doc.paragraphs[115],
        "How Hashing Protects Citizen Privacy: A cryptographic hash works like a fingerprint scanner. When a chassis number or plate number is "
        "hashed, it produces a unique 32-byte code. A traffic police officer can easily check if a car's plate matches the blockchain record by "
        "hashing the plate and checking for a match. However, no outsider browsing the blockchain can reverse the hash to find out who the owner "
        "is or where they live.",
        size=11.5
    )
    set_para_text(
        doc.paragraphs[117],
        "Clear Separation of RTO Roles: To prevent corruption, powers are strictly separated. The RTO Registrar can register new vehicles and "
        "approve ownership transfers, but cannot scrap vehicles. The RTO Enforcement Officer can mark damaged vehicles as Scrapped, but cannot "
        "register or transfer vehicles. Most importantly, an automated smart contract rule strictly bars any RTO or Admin wallet from ever "
        "owning a vehicle title, preventing transport officials from transferring impounded cars to themselves.",
        size=11.5
    )
    set_para_text(
        doc.paragraphs[119],
        "Two-Step Transfer & Permanent Scrappage: When selling a car, the citizen initiates a transfer request in the portal (Form 29). "
        "The RTO Registrar checks that road taxes are paid, no police cases exist, and the bank loan is cleared before officially approving "
        "the transfer (Form 30). Once approved, the title moves to the buyer and the complete history is saved. If an officer marks a vehicle "
        "as Scrapped, it is permanently locked on the blockchain and can never be re-registered or transferred again.",
        size=11.5
    )

    # Implementation (Paragraphs 123, 124, 125, 126)
    set_para_text(
        doc.paragraphs[123],
        "The system was built using two main parts: the Solidity smart contract and the React frontend web application.",
        size=11.5, bold=True
    )
    set_para_text(
        doc.paragraphs[124],
        "1. Smart Contract (VehicleRegistry.sol): Written in Solidity 0.8.24 and deployed on a local Ethereum network (Hardhat). "
        "It uses OpenZeppelin standards: ERC-721 for vehicle title NFTs, AccessControl for managing RTO roles, and Pausable for emergency controls. "
        "The main functions are: registerVehicle() to mint new titles, transferOwnershipByRTO() to approve and execute transfers, "
        "initiateTransfer() for citizens to request a transfer, and updateVehicleState() for officers to scrap vehicles.",
        size=11.5
    )
    set_para_text(
        doc.paragraphs[125],
        "2. Anti-Corruption Security Rule (_update hook): We customized OpenZeppelin's internal transfer function so that every time a token "
        "transfer is attempted, the contract checks if the recipient holds an RTO or Admin role. If an RTO address is detected, the transaction "
        "is immediately rejected with the error: PrivilegedAddressCannotOwnVehicle. This rule is enforced by blockchain code and cannot be bypassed.",
        size=11.5
    )
    set_para_text(
        doc.paragraphs[126],
        "3. Web Application: Built using React 18, Vite, and Ethers.js v6. It provides four simple portals: a Public Search Portal (to check if a "
        "car is Active or Scrapped), a Citizen Dashboard (to view verified digital HSRP number plates and transfer history), an RTO Registrar "
        "Desk (to register vehicles and approve transfers), and an RTO Officer Desk (to inspect and scrap vehicles).",
        size=11.5
    )
    for idx in range(127, 129):
        doc.paragraphs[idx].text = ""

    # Output of Mini Project (Paragraphs 130, 131)
    set_para_text(
        doc.paragraphs[130],
        "The project was thoroughly tested using an automated test suite written in Hardhat and Chai. A total of 40 test cases were executed, "
        "and all 40 passed with a 100% success rate in 5 seconds. All 16 negative security attack scenarios were successfully blocked:",
        size=11.5
    )
    set_para_text(
        doc.paragraphs[131],
        "• Unauthorized citizens cannot register vehicles or modify legal status.\n"
        "• Duplicate chassis number (VIN) and number plate registrations are rejected on-chain.\n"
        "• RTO authority wallets are strictly blocked from receiving or owning vehicle titles.\n"
        "• Scrapped vehicles cannot be transferred, sold, or re-registered under any condition.\n"
        "• Full ownership history is permanently recorded with accurate timestamps for every legal transfer.\n"
        "• Emergency pause controls successfully block all operations during administrative lockouts.",
        size=11
    )
    for idx in range(132, 134):
        doc.paragraphs[idx].text = ""

    # Skill Developed (Paragraph 137)
    set_para_text(
        doc.paragraphs[137],
        "Working on this project gave us valuable hands-on experience in blockchain and cybersecurity. We learned how to write real, deployable "
        "smart contracts in Solidity and understood how Ethereum automatically enforces rules without any middleman. We practically applied "
        "Keccak-256 cryptographic hashing to solve the privacy challenge of public blockchains, keeping citizen data safe. We also learned how to "
        "design Role-Based Access Control so that government officials have clear, limited powers. Finally, we learned test-driven development "
        "by writing 40 automated tests to find and fix security vulnerabilities before deploying the system.",
        size=11.5
    )

    # Applications (Paragraph 140)
    set_para_text(
        doc.paragraphs[140],
        "This system can be used in several real-world areas: (1) State RTO Modernization — Transport departments can use it to completely stop "
        "fake RC books, stolen vehicle cloning, and illegal duplicate registrations across states; (2) Used-Car Marketplaces (Cars24, Spinny) — "
        "Buyers can instantly verify the complete ownership history and legal standing of a second-hand car without manual paperwork; "
        "(3) Insurance Companies — When a car is declared a total loss in an accident, the insurer can permanently mark it as scrapped on-chain "
        "so salvage cars cannot be rebuilt and resold; and (4) Bank Auto Loans — Banks can place digital hypothecation liens on vehicle titles "
        "until car loans are fully repaid.",
        size=11.5
    )

    # Area of Future Improvement (Paragraph 147)
    set_para_text(
        doc.paragraphs[147],
        "In the future, this system can be expanded with several advanced features: (1) Zero-Knowledge Proofs (zk-SNARKs) — Allowing citizens "
        "to prove to traffic police that their car is legally registered without revealing their wallet address or identity; (2) Decentralized "
        "Storage (IPFS) — Storing vehicle inspection photos, fitness certificates, and pollution documents on IPFS instead of central servers; "
        "(3) IoT and FASTag Integration — Connecting vehicle OBD-II chips and FASTag scanners to automatically record tamper-proof odometer "
        "readings on the blockchain, eliminating odometer fraud; and (4) Layer-2 Networks — Deploying on Polygon or Arbitrum to reduce transaction "
        "costs to fractions of a rupee for nationwide use.",
        size=11.5
    )
    for idx in range(148, 150):
        doc.paragraphs[idx].text = ""

    # Conclusion (Paragraph 154 heading & Paragraph 155 body)
    p154 = doc.paragraphs[154]
    p154.text = "Conclusion"
    p154.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r154 = p154.runs[0]
    r154.font.name = FONT_NAME
    r154.font.size = Pt(12)
    r154.font.bold = True

    if len(doc.paragraphs) > 155:
        p155 = doc.paragraphs[155]
    else:
        p155 = doc.add_paragraph()

    set_para_text(
        p155,
        "This mini-project demonstrates that blockchain technology can solve real, everyday problems in vehicle registration and used-car buying. "
        "By replacing centralized databases with immutable smart contracts, we created a system where records cannot be secretly altered, "
        "citizen privacy is protected through cryptographic hashing, and salvage car fraud is permanently stopped. The strict separation "
        "between the RTO Registrar and the RTO Officer ensures that no single person has unlimited authority. With all 40 automated tests "
        "passing at 100%, the project shows that code-enforced rules provide real security, transparency, and trust for modern transport governance.",
        size=11.5
    )

    # ── 6. Save modified document ──────────────────────────────────────
    doc.save(OUTPUT_PROJECT)
    print("Saved to project:", OUTPUT_PROJECT, f"({os.path.getsize(OUTPUT_PROJECT)//1024} KB)")

    for path, name in [(OUTPUT_DESKTOP, "Desktop"), (OUTPUT_DOWNLOADS, "Downloads"), (OUTPUT_ARTIFACT, "Artifacts")]:
        try:
            shutil.copy2(OUTPUT_PROJECT, path)
            print(f"Copied to {name}:", path)
        except Exception as e:
            print(f"{name} copy failed:", e)

    return OUTPUT_PROJECT

if __name__ == "__main__":
    build()
