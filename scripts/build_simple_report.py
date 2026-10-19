import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build():
    doc = docx.Document()

    # ── Page setup ──────────────────────────────────────────────────────
    sec = doc.sections[0]
    sec.page_width  = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin = sec.bottom_margin = Inches(1.0)
    sec.left_margin = sec.right_margin = Inches(1.1)
    sec.different_first_page_header_footer = True

    # Running header (non-first pages)
    hdr = sec.header
    ph = hdr.paragraphs[0]
    ph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rh = ph.add_run("MGM University  |  Mini Project Report  |  AY 2025–26")
    rh.font.name, rh.font.size = "Calibri", Pt(8.5)
    rh.font.color.rgb = RGBColor(130,130,130)
    rh.font.italic = True

    # Running footer with page number
    ftr = sec.footer
    pf = ftr.paragraphs[0]
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = pf.add_run("Decentralized Vehicle Registration System   |   Page ")
    rf.font.name, rf.font.size = "Calibri", Pt(9)
    rf.font.color.rgb = RGBColor(110,110,110)
    pf._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>'))

    # ── Colour shortcuts ─────────────────────────────────────────────────
    NAVY  = RGBColor(31,73,125)
    SLATE = RGBColor(52,100,150)
    BODY  = RGBColor(40,40,40)
    MUTED = RGBColor(100,100,100)
    WHITE = RGBColor(255,255,255)
    HEX_NAVY  = "1F497D"
    HEX_ROW   = "F0F4F9"
    HEX_BORDER= "B0BEC5"

    # ── Helpers ──────────────────────────────────────────────────────────
    def shading(cell, hex_col):
        cell._tc.get_or_add_tcPr().append(
            parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_col}"/>'))

    def cell_pad(cell, v=100, h=120):
        cell._tc.get_or_add_tcPr().append(parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{v}" w:type="dxa"/><w:bottom w:w="{v}" w:type="dxa"/>'
            f'<w:left w:w="{h}" w:type="dxa"/><w:right w:w="{h}" w:type="dxa"/>'
            f'</w:tcMar>'))

    def tbl_border(tbl, col=HEX_BORDER):
        tbl._tbl.tblPr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:color="{col}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:color="{col}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:color="{col}"/>'
            f'<w:insideV w:val="single" w:sz="4" w:color="{col}"/>'
            f'<w:left w:val="none"/><w:right w:val="none"/>'
            f'</w:tblBorders>'))

    def hdr_row(row, widths=None):
        for i, cell in enumerate(row.cells):
            shading(cell, HEX_NAVY)
            cell_pad(cell, 130, 120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths: cell.width = widths[i]
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name, r.font.size = "Calibri", Pt(10)
                    r.font.bold, r.font.color.rgb = True, WHITE

    def data_row(row, alt=False, widths=None, aligns=None):
        for i, cell in enumerate(row.cells):
            if alt: shading(cell, HEX_ROW)
            cell_pad(cell, 90, 120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths: cell.width = widths[i]
            for p in cell.paragraphs:
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
                if aligns: p.alignment = aligns[i]
                for r in p.runs:
                    r.font.name, r.font.size, r.font.color.rgb = "Calibri", Pt(9.5), BODY

    def h1(text):
        p = doc.add_paragraph(style="Heading 1")
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after  = Pt(8)
        r = p.add_run(text)
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Calibri", Pt(14), True, NAVY
        return p

    def h2(text):
        p = doc.add_paragraph(style="Heading 2")
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after  = Pt(5)
        r = p.add_run(text)
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Calibri", Pt(11.5), True, SLATE
        return p

    def body(text, bold_lead=""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(6)
        p.paragraph_format.line_spacing = 1.2
        if bold_lead:
            bl = p.add_run(bold_lead)
            bl.font.name, bl.font.size, bl.font.bold, bl.font.color.rgb = "Calibri", Pt(10.5), True, BODY
        r = p.add_run(text)
        r.font.name, r.font.size, r.font.color.rgb = "Calibri", Pt(10.5), BODY
        return p

    def bullet(text, lead=""):
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(4)
        if lead:
            bl = p.add_run(lead)
            bl.font.name, bl.font.size, bl.font.bold, bl.font.color.rgb = "Calibri", Pt(10.5), True, BODY
        r = p.add_run(text)
        r.font.name, r.font.size, r.font.color.rgb = "Calibri", Pt(10.5), BODY
        return p

    def note_box(text, label="NOTE"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0,0)
        cell.width = Inches(6.3)
        shading(cell, "EBF3FB")
        cell_pad(cell, 130, 160)
        cell._tc.get_or_add_tcPr().append(parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="18" w:color="{HEX_NAVY}"/>'
            f'<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
            f'</w:tcBorders>'))
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        r_l = p.add_run(f"{label}: ")
        r_l.font.name, r_l.font.size, r_l.font.bold, r_l.font.color.rgb = "Calibri", Pt(10), True, NAVY
        r_t = p.add_run(text)
        r_t.font.name, r_t.font.size, r_t.font.italic, r_t.font.color.rgb = "Calibri", Pt(9.5), True, BODY
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def code_box(text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0,0)
        cell.width = Inches(6.3)
        shading(cell, "F4F6F9")
        cell_pad(cell, 110, 140)
        cell._tc.get_or_add_tcPr().append(parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:color="{HEX_BORDER}"/>'
            f'<w:left w:val="single" w:sz="4" w:color="{HEX_BORDER}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:color="{HEX_BORDER}"/>'
            f'<w:right w:val="single" w:sz="4" w:color="{HEX_BORDER}"/>'
            f'</w:tcBorders>'))
        for i, line in enumerate(text.strip().split("\n")):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(line)
            r.font.name, r.font.size, r.font.color.rgb = "Consolas", Pt(8.5), RGBColor(25,40,60)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def divider():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(6)
        r = p.add_run("─" * 60)
        r.font.size, r.font.color.rgb = Pt(9), RGBColor(180,180,180)

    # ====================================================================
    # COVER PAGE
    # ====================================================================
    def cp(text, size, bold=False, color=NAVY, space_before=0, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after  = Pt(space_after)
        r = p.add_run(text)
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Calibri", Pt(size), bold, color
        return p

    cp("School of Engineering and Technology", 16, bold=True, color=NAVY, space_before=18, space_after=2)
    cp("MGM University, Chh. Sambhajinagar", 13, bold=True, color=SLATE, space_after=2)
    cp("Academic Year 2025 – 2026", 11, color=MUTED, space_after=10)
    divider()
    cp("Mini Project Report", 13, bold=True, color=SLATE, space_before=6, space_after=4)
    cp("On", 11, color=MUTED, space_after=8)
    cp('"Decentralized Vehicle Registration &\nImmutable Title Management System\nUsing Blockchain Technology"',
       18, bold=True, color=NAVY, space_after=14)
    cp("Submitted in partial fulfilment for the Sixth Semester of B.Tech (Integrated) in", 10.5, color=BODY, space_after=2)
    cp("COMPUTER SCIENCE AND ENGINEERING", 13, bold=True, color=NAVY, space_after=2)
    cp("(IoT, Cyber Security Including Blockchain Technology)", 11, bold=True, color=SLATE, space_after=14)
    cp("Submitted By", 11, bold=True, color=MUTED, space_after=6)

    tbl = doc.add_table(rows=3, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl)
    ws = [Inches(3.3), Inches(2.4)]
    tbl.rows[0].cells[0].paragraphs[0].text = "Name of Student"
    tbl.rows[0].cells[1].paragraphs[0].text = "PRN Number"
    hdr_row(tbl.rows[0], ws)
    for rw, nm, prn, alt in [
        (tbl.rows[1], "Rohit Dhananjay Kale",   "202356112007", False),
        (tbl.rows[2], "Soham Bhagwan Deshmukh", "202356112008", True)]:
        rw.cells[0].paragraphs[0].text = nm
        rw.cells[1].paragraphs[0].text = prn
        data_row(rw, alt, ws, [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])

    cp("Under the Guidance of", 10.5, bold=True, color=MUTED, space_before=18, space_after=2)
    cp("Ms. Bhagat C. B.", 12, bold=True, color=NAVY, space_after=2)
    cp("(Assistant Professor)  |  Department of Computer Science and Engineering", 10, color=BODY)
    doc.add_page_break()

    # ====================================================================
    # CERTIFICATE
    # ====================================================================
    cp("CERTIFICATE", 18, bold=True, color=NAVY, space_before=22, space_after=14)

    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert.paragraph_format.space_before = Pt(8)
    p_cert.paragraph_format.space_after  = Pt(8)
    p_cert.paragraph_format.line_spacing = 1.3
    def cr(t, bold=False):
        rr = p_cert.add_run(t)
        rr.font.name, rr.font.size, rr.font.bold = "Calibri", Pt(11), bold
    cr("This is to certify that ")
    cr("Rohit Dhananjay Kale", True); cr(" and ")
    cr("Soham Bhagwan Deshmukh", True)
    cr(" with PRN Numbers ")
    cr("202356112007", True); cr(" and ")
    cr("202356112008", True)
    cr(" have successfully completed their Mini-Project entitled:\n\n")

    cp('"Decentralized Vehicle Registration & Immutable Title Management System Using Blockchain Technology"',
       13, bold=True, color=NAVY, space_after=12)

    p_tail = doc.add_paragraph()
    p_tail.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_tail.paragraph_format.space_after = Pt(38)
    p_tail.paragraph_format.line_spacing = 1.3
    rt = p_tail.add_run(
        "in the Course of Blockchain Technology & Cyber Security in the Sixth Semester, "
        "during the tenure of completing the B.Tech (Integrated) programme in Computer Science "
        "and Engineering (IoT, Cyber Security Including Blockchain Technology) from School of "
        "Engineering and Technology, MGM University, Chh. Sambhajinagar.")
    rt.font.name, rt.font.size = "Calibri", Pt(11)

    tbl_s = doc.add_table(rows=2, cols=2)
    tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    sw = [Inches(3.3), Inches(3.2)]
    tbl_s.rows[0].cells[0].paragraphs[0].text = "______________________________\nMs. Bhagat C. B."
    tbl_s.rows[0].cells[1].paragraphs[0].text = "______________________________\nMs. Rathi S. R."
    tbl_s.rows[1].cells[0].paragraphs[0].text = "Guide, Asst. Professor\nDept. of CSE\nSchool of Engineering & Technology\nMGM University, Chh. Sambhajinagar"
    tbl_s.rows[1].cells[1].paragraphs[0].text = "Head of Department (HOD)\nDept. of CSE\nSchool of Engineering & Technology\nMGM University, Chh. Sambhajinagar"
    for ri, rw in enumerate(tbl_s.rows):
        for cell in rw.cells:
            cell.width = sw[0] if cell == rw.cells[0] else sw[1]
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.2
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(11 if ri==0 else 9.5)
                    r.font.bold = (ri == 0)
                    r.font.color.rgb = NAVY if ri==0 else BODY
    doc.add_page_break()

    # ====================================================================
    # ACKNOWLEDGEMENT
    # ====================================================================
    cp("ACKNOWLEDGEMENT", 16, bold=True, color=NAVY, space_before=20, space_after=14)
    body("We are grateful to our project guide, Ms. Bhagat C. B. (Assistant Professor, Dept. of CSE), "
         "for her continuous support and guidance throughout this project. Her suggestions helped us "
         "understand both the technical and practical aspects of blockchain development.")
    body("We also thank Ms. Rathi S. R. (Head of Department, Dept. of CSE) and the management of "
         "MGM University, Chh. Sambhajinagar, for providing excellent resources and an environment "
         "that encouraged us to work on real-world problems.")
    body("Finally, we thank our family and friends for their constant encouragement and motivation.")
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(24)
    rs = p_sign.add_run("Rohit Dhananjay Kale  (PRN: 202356112007)\n"
                         "Soham Bhagwan Deshmukh  (PRN: 202356112008)\n"
                         "B.Tech (Integrated) CSE — Semester VI")
    rs.font.name, rs.font.size, rs.font.bold, rs.font.color.rgb = "Calibri", Pt(10), True, NAVY
    doc.add_page_break()

    # ====================================================================
    # TABLE OF CONTENTS
    # ====================================================================
    cp("TABLE OF CONTENTS", 15, bold=True, color=NAVY, space_before=14, space_after=12)
    toc_data = [
        ("1","Abstract","01"), ("2","Introduction","02"), ("3","Course Outcomes Integrated","03"),
        ("4","Action Plan","04"), ("5","Literature Review","05"),
        ("6","Proposed Methodology & System Architecture","06"),
        ("7","Implementation Details","09"), ("8","Project Outputs & Test Results","11"),
        ("9","Skills Developed","13"), ("10","Applications","14"),
        ("11","Future Improvements","15"), ("12","Conclusion","16")
    ]
    tbl_toc = doc.add_table(rows=len(toc_data)+1, cols=3)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_toc)
    tw = [Inches(0.9), Inches(4.7), Inches(0.9)]
    tbl_toc.rows[0].cells[0].paragraphs[0].text = "Sr. No."
    tbl_toc.rows[0].cells[1].paragraphs[0].text = "Chapter / Section"
    tbl_toc.rows[0].cells[2].paragraphs[0].text = "Page No."
    hdr_row(tbl_toc.rows[0], tw)
    for i, (n, t, pg) in enumerate(toc_data):
        r = tbl_toc.rows[i+1]
        r.cells[0].paragraphs[0].text = n
        r.cells[1].paragraphs[0].text = t
        r.cells[2].paragraphs[0].text = pg
        data_row(r, i%2==1, tw, [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])
    doc.add_page_break()

    # ====================================================================
    # 1. ABSTRACT
    # ====================================================================
    h1("1. Abstract")
    body("Vehicle registration in India is currently managed by Regional Transport Offices (RTOs) using centralized "
         "government databases. These systems have several serious problems — a single database failure can bring the "
         "entire system down, corrupt officials can quietly edit records, and citizens have no way to independently verify "
         "a car's ownership history. Buyers also frequently fall victim to 'ghost vehicles' — stolen cars sold using fake "
         "Registration Certificates.")
    body("This mini-project builds a Decentralized Vehicle Registration System using the Ethereum blockchain. In this "
         "system, every vehicle title is represented as a digital token (using the ERC-721 standard). No single person or "
         "server controls the records. Once a vehicle is registered or transferred, the record is permanently stored on the "
         "blockchain and cannot be changed by anyone — not even the developers.")
    body("To protect citizen privacy, we do not store personal details like names or phone numbers on the public blockchain. "
         "Instead, we store only a 'fingerprint' (Keccak-256 hash) of the vehicle's chassis number and plate. The actual "
         "personal data stays in a private, encrypted storage vault.")
    body("The system separates government duties clearly: an RTO Registrar registers vehicles and approves ownership "
         "transfers, while an RTO Enforcement Officer handles vehicle scrappage. Neither role can interfere with the other's "
         "work. The project was validated with 40 automated tests, all of which passed successfully.")
    note_box("This system replaces trust in a person or institution with trust in blockchain code — making vehicle records "
             "tamper-proof, verifiable, and private at the same time.", "KEY IDEA")
    doc.add_page_break()

    # ====================================================================
    # 2. INTRODUCTION
    # ====================================================================
    h1("2. Introduction")
    body("In India, every vehicle on the road must be registered with the local Regional Transport Office (RTO). The "
         "Registration Certificate (RC) is the proof of ownership. If you want to sell a car, buy a second-hand car, "
         "or insure a vehicle, the RC is the most important document involved.")
    body("Today, RTOs maintain these records on centralized computer systems (like the national Vahan portal). While "
         "this is better than old paper-based records, it still has many weaknesses that cause real harm to ordinary citizens.")

    h2("2.1 Problems with the Current System")
    bullet("Criminals sometimes steal a car, duplicate its chassis plate, and obtain fake RCs. Innocent buyers "
           "purchase these vehicles and later lose them to law enforcement with no refund.",
           "Fake Vehicles & Forged Documents: ")
    bullet("All records live on a central server. If a corrupt official or hacker gains access, they can change the "
           "owner's name, remove a bank loan record, or alter a vehicle's status. There is no permanent audit trail.",
           "Database Tampering: ")
    bullet("When a car buyer wants to check the history of a second-hand vehicle, they must rely on paperwork "
           "provided by the seller. There is no independent, public way to verify the vehicle's complete ownership history.",
           "No Transparent History: ")
    bullet("When a car is totaled in an accident and declared a 'total loss' by insurance, the chassis is sometimes "
           "secretly rebuilt and resold. Buyers have no way of knowing this happened.",
           "Salvage Car Fraud: ")
    bullet("Government administrative staff sometimes use their access to transfer impounded or seized vehicles to "
           "personal accounts. There are no automated technical checks to stop this.",
           "Authority Abuse: ")

    h2("2.2 How This Project Solves These Problems")
    bullet("Vehicle titles are stored as blockchain tokens (ERC-721 NFTs). No record can ever be changed or deleted once written.", "Tamper-Proof Records: ")
    bullet("Only Keccak-256 fingerprints of VINs and plates are stored on-chain. Citizen names and addresses stay in private encrypted storage.", "Privacy by Design: ")
    bullet("Every change of ownership is permanently recorded on the blockchain with a timestamp, creating a complete, verifiable history.", "Full Ownership History: ")
    bullet("RTO Registrars can register and transfer. RTO Officers can only scrap. Neither can do the other's job. Authorities cannot own vehicles.", "Strict Role Separation: ")
    bullet("Once an officer marks a vehicle as 'Scrapped', no one can ever transfer or re-register it again. The rule is enforced by code.", "Permanent Scrappage Lock: ")
    doc.add_page_break()

    # ====================================================================
    # 3. COURSE OUTCOMES INTEGRATED
    # ====================================================================
    h1("3. Course Outcomes Integrated")
    body("This project covers the following Course Outcomes (COs) from the B.Tech curriculum in Computer Science "
         "Engineering (IoT, Cyber Security & Blockchain Technology):")

    co_data = [
        ("CO1","Cryptographic Foundations",
         "Used Keccak-256 hashing to protect vehicle identity data and EIP-191 digital signatures for wallet login."),
        ("CO2","Smart Contract Development",
         "Wrote a complete Solidity 0.8.24 smart contract using OpenZeppelin libraries for ERC-721 tokens, access control, and pausable safety."),
        ("CO3","Role-Based Access Control (RBAC)",
         "Designed separate roles — Admin, RTO Registrar, RTO Officer, and Citizen — each with strict, non-overlapping permissions enforced on the blockchain."),
        ("CO4","Decentralized Application (DApp) Development",
         "Built a full React 18 + Vite web application that communicates with the Ethereum node through Ethers.js v6. Includes route guards and live transaction handling."),
        ("CO5","Security Testing & Verification",
         "Wrote 40 automated Hardhat test cases covering 16 negative security scenarios (unauthorized access attempts, privilege escalations, and state machine violations).")
    ]
    tbl_co = doc.add_table(rows=len(co_data)+1, cols=3)
    tbl_co.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_co)
    cw = [Inches(0.9), Inches(1.9), Inches(3.7)]
    tbl_co.rows[0].cells[0].paragraphs[0].text = "CO"
    tbl_co.rows[0].cells[1].paragraphs[0].text = "Topic"
    tbl_co.rows[0].cells[2].paragraphs[0].text = "How this Project Uses It"
    hdr_row(tbl_co.rows[0], cw)
    for i, (co, topic, how) in enumerate(co_data):
        r = tbl_co.rows[i+1]
        r.cells[0].paragraphs[0].text = co
        r.cells[1].paragraphs[0].text = topic
        r.cells[2].paragraphs[0].text = how
        data_row(r, i%2==1, cw, [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    doc.add_page_break()

    # ====================================================================
    # 4. ACTION PLAN
    # ====================================================================
    h1("4. Action Plan")
    body("The project was completed in six phases over three months (August to October 2025). Both team members "
         "divided the work clearly based on their strengths:")
    act_data = [
        ("1","Study of RTO workflows (Form 29/30) & System Architecture Design","01-Aug-2025","12-Aug-2025","Both Members"),
        ("2","Smart Contract Development: VehicleRegistry.sol (ERC-721 & RBAC)","13-Aug-2025","28-Aug-2025","Rohit Kale"),
        ("3","Privacy Layer: Keccak-256 Commitments & Off-Chain Encrypted Storage","29-Aug-2025","10-Sep-2025","Soham Deshmukh"),
        ("4","Security Boundaries: Transfer Workflow & Scrappage State Machine","11-Sep-2025","24-Sep-2025","Rohit Kale"),
        ("5","Frontend Development: React/Vite UI, Citizen & RTO Portals","25-Sep-2025","12-Oct-2025","Soham Deshmukh"),
        ("6","Testing (40 Test Cases), Live Blockchain Proof & Report","13-Oct-2025","28-Oct-2025","Both Members")
    ]
    tbl_a = doc.add_table(rows=len(act_data)+1, cols=5)
    tbl_a.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_a)
    aw = [Inches(0.55), Inches(2.85), Inches(1.05), Inches(1.05), Inches(0.98)]
    for ci, hd in enumerate(["No.","Activity","Start Date","End Date","Responsible"]):
        tbl_a.rows[0].cells[ci].paragraphs[0].text = hd
    hdr_row(tbl_a.rows[0], aw)
    for i, (n, act, sd, ed, who) in enumerate(act_data):
        rw = tbl_a.rows[i+1]
        for ci, val in enumerate([n, act, sd, ed, who]):
            rw.cells[ci].paragraphs[0].text = val
        data_row(rw, i%2==1, aw,
                 [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    doc.add_page_break()

    # ====================================================================
    # 5. LITERATURE REVIEW
    # ====================================================================
    h1("5. Literature Review")
    body("We studied existing systems and earlier research before designing our solution.")

    h2("5.1 Current Indian System: Vahan / Parivahan")
    body("The Government of India's Vahan portal is the national vehicle registration database used by all RTOs. "
         "It improved record-keeping compared to manual registers and enabled online tax payments. However, "
         "since all data lives on central government servers, it remains vulnerable to cyber attacks, server failures, "
         "and insider misuse. When a record is altered, there is no automatic way to detect or reverse it.")

    h2("5.2 Earlier Blockchain Research")
    body("Several university projects have attempted to use blockchain for vehicle registration. However, most had two common flaws:")
    bullet("They stored citizens' real names, phone numbers, and addresses directly on the public blockchain. "
           "Since anyone can read blockchain data, this is a serious privacy violation.",
           "Privacy Violation: ")
    bullet("They allowed citizens to transfer vehicle ownership directly like sending a WhatsApp message, "
           "bypassing all legal checks. In India, this is illegal — RTO approval is mandatory for transfers.",
           "Ignoring Legal Requirements: ")

    h2("5.3 Comparison: Old vs New vs Our Approach")
    comp = [
        ("Tamper-Proof Records","No — can be edited by admins","Partially — depends on design","Yes — blockchain is immutable"),
        ("Citizen Privacy Protection","Moderate — central server risk","No — PII exposed on public chain","Yes — only hashes stored on-chain"),
        ("Transparent Ownership History","No — not publicly accessible","Partially — basic records","Yes — complete history on blockchain"),
        ("Legal Transfer Compliance","Manual paperwork, slow","Not enforced","Yes — RTO Registrar must approve"),
        ("Role Separation (RTO)","Loose — admin has full access","None — single admin","Strict — Registrar vs Officer vs Admin"),
        ("Scrappage Enforcement","Paper certificate, not enforced","Boolean flag — easily bypassed","Permanent state lock — cannot be undone"),
    ]
    tbl_c = doc.add_table(rows=len(comp)+1, cols=4)
    tbl_c.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_c)
    kcw = [Inches(1.55), Inches(1.5), Inches(1.5), Inches(1.95)]
    for ci, hd in enumerate(["Feature","Traditional Database","Earlier Blockchain","Our System"]):
        tbl_c.rows[0].cells[ci].paragraphs[0].text = hd
    hdr_row(tbl_c.rows[0], kcw)
    for i, row_data in enumerate(comp):
        rw = tbl_c.rows[i+1]
        for ci, val in enumerate(row_data):
            rw.cells[ci].paragraphs[0].text = val
        data_row(rw, i%2==1, kcw)
    doc.add_page_break()

    # ====================================================================
    # 6. PROPOSED METHODOLOGY & SYSTEM ARCHITECTURE
    # ====================================================================
    h1("6. Proposed Methodology & System Architecture")
    body("Our system is built in four layers that work together. Think of it like a building — "
         "users interact at the front desk, security checks happen at the entrance, data is converted "
         "to protected form in the middle, and the final records are locked permanently in the vault.")

    h2("6.1 The Four Layers")
    bullet("Ordinary citizens, used-car buyers, RTO Registrars, and RTO Officers. Each has a different portal in the web app.", "Layer 1 — Users: ")
    bullet("A React 18 web application (built with Vite). Each portal has a security guard (CitizenGuard or RtoGuard) "
           "that blocks the wrong type of user from entering. If you switch your wallet in MetaMask, you are immediately logged out.",
           "Layer 2 — Web App & Security: ")
    bullet("Before any data is sent to the blockchain, sensitive identifiers (VIN and plate number) are converted "
           "into Keccak-256 fingerprint hashes. Names, addresses, and phone numbers stay in a private encrypted storage.",
           "Layer 3 — Privacy Conversion: ")
    bullet("The Ethereum smart contract (VehicleRegistry.sol). This is where vehicle titles are created, ownership is "
           "transferred, vehicles are scrapped, and history is permanently recorded.",
           "Layer 4 — Blockchain (Smart Contract): ")

    h2("6.2 How Hashing Protects Privacy")
    body("A Keccak-256 hash works like a fingerprint scanner. When you scan a fingerprint, you get a unique code. "
         "You can check if two fingerprints match — but you can never rebuild the original finger from the code alone. "
         "We apply the same idea to vehicle numbers:")
    bullet("The chassis number (e.g., MA3ERLF3S00123456) is hashed → a 32-byte code is stored on blockchain.", "VIN Hash: ")
    bullet("The number plate (e.g., MH20DE1234) is hashed → a 32-byte code is stored on blockchain.", "Plate Hash: ")
    body("This means a police officer or RTO official can verify a vehicle is genuine by entering the plate number — "
         "the system hashes it and checks if it matches the stored fingerprint. But no one browsing the blockchain "
         "can reverse the hash to find out who owns the car.")

    h2("6.3 The Two RTO Roles — Why Separate?")
    body("In a traditional system, one senior admin often has all powers — registering cars, transferring them, "
         "and scrapping them. This is dangerous because a single corrupt official can misuse everything. "
         "We split the work into two clearly different roles:")

    tbl_r = doc.add_table(rows=3, cols=3)
    tbl_r.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_r)
    rw2 = [Inches(1.6), Inches(2.2), Inches(2.7)]
    for ci, hd in enumerate(["Role","What They CAN Do","What They CANNOT Do"]):
        tbl_r.rows[0].cells[ci].paragraphs[0].text = hd
    hdr_row(tbl_r.rows[0], rw2)
    tbl_r.rows[1].cells[0].paragraphs[0].text = "RTO Registrar"
    tbl_r.rows[1].cells[1].paragraphs[0].text = "Register new vehicles\nApprove ownership transfers"
    tbl_r.rows[1].cells[2].paragraphs[0].text = "Cannot scrap vehicles\nCannot own vehicle titles"
    tbl_r.rows[2].cells[0].paragraphs[0].text = "RTO Enforcement Officer"
    tbl_r.rows[2].cells[1].paragraphs[0].text = "Mark vehicles as Scrapped\nConduct fitness inspections"
    tbl_r.rows[2].cells[2].paragraphs[0].text = "Cannot register vehicles\nCannot transfer titles\nCannot own vehicle titles"
    data_row(tbl_r.rows[1], False, rw2)
    data_row(tbl_r.rows[2], True, rw2)

    body("")
    note_box("ANTI-CORRUPTION RULE: Neither the Admin nor any RTO role holder can ever own a vehicle title. "
             "If anyone tries to transfer a car to an RTO wallet, the smart contract automatically rejects it "
             "with the error: PrivilegedAddressCannotOwnVehicle.", "IMPORTANT")

    h2("6.4 How Ownership Transfer Works (Step by Step)")
    body("Transferring a vehicle title requires two steps — one by the citizen, one by the RTO. This mirrors "
         "the legal Form 29 (Notice of Transfer) and Form 30 (Application for Transfer) process:")
    bullet("The current vehicle owner logs into the Citizen Portal and submits a transfer request "
           "with the buyer's wallet address. The car stays in the seller's name at this point.",
           "Step 1 — Citizen Initiates: ")
    bullet("The RTO Registrar reviews the request, verifies that all legal conditions are met "
           "(taxes paid, no fines pending, bank loan cleared), and calls transferOwnershipByRTO() on the smart contract.",
           "Step 2 — RTO Registrar Approves: ")
    bullet("The smart contract moves the digital title token to the buyer's wallet and permanently adds "
           "the buyer to the vehicle's ownership history with a timestamp.",
           "Step 3 — Title Moves, History Updated: ")

    h2("6.5 Vehicle Scrappage (Permanent Terminal State)")
    body("When a vehicle is totaled in a crash or reaches 15 years of age, an RTO Enforcement Officer can "
         "mark it as 'Scrapped'. This is a permanent, irreversible action on the blockchain:")
    bullet("All pending transfer requests are immediately cancelled.", "")
    bullet("The vehicle title is permanently locked — no transfers, no re-registrations, ever.", "")
    bullet("Any future attempt to move the token reverts with the error VehicleIsScrapped.", "")
    body("This completely prevents the common fraud of rebuilding crashed cars and selling them to unknowing buyers.")
    doc.add_page_break()

    # ====================================================================
    # 7. IMPLEMENTATION DETAILS
    # ====================================================================
    h1("7. Implementation Details")
    body("The project has two main parts: the smart contract running on Ethereum, and the web application "
         "that citizens and RTO staff use to interact with it.")

    h2("7.1 Smart Contract: VehicleRegistry.sol")
    body("Written in Solidity 0.8.24, deployed on a local Ethereum node (Hardhat). "
         "Built using trusted OpenZeppelin v5.0 library contracts:")
    bullet("ERC721 — for creating non-fungible vehicle title tokens.")
    bullet("AccessControl — for managing Admin, Registrar, and Officer roles.")
    bullet("Pausable — for emergency shutdowns by the Admin.")
    body("The main data structures stored for each vehicle:")
    code_box("""struct Vehicle {
    bytes32 vehicleId;              // Unique ID (hash of VIN + Plate)
    bytes32 vinCommitment;          // Keccak-256 hash of chassis number
    bytes32 registrationCommitment; // Keccak-256 hash of number plate
    VehicleState state;             // Active (0)  or  Scrapped (1)
    address titleHolder;            // Current owner's wallet address
    bytes32 metadataHash;           // Hash of vehicle specs (make, model, year)
    uint64 registeredAt;            // When it was first registered
}

struct OwnershipRecord {
    address owner;                  // Owner's wallet address
    uint64 transferredAt;           // When they received the title
}""")

    h2("7.2 Key Smart Contract Functions")
    tbl_f = doc.add_table(rows=5, cols=3)
    tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_f)
    fw = [Inches(1.7), Inches(1.3), Inches(3.5)]
    for ci, hd in enumerate(["Function","Who Can Call It","What It Does"]):
        tbl_f.rows[0].cells[ci].paragraphs[0].text = hd
    hdr_row(tbl_f.rows[0], fw)
    funcs = [
        ("registerVehicle(...)","RTO Registrar Only","Creates a new vehicle record and mints an ERC-721 title token for the citizen."),
        ("transferOwnershipByRTO(...)","RTO Registrar Only","Moves the title token to the new owner and records the transfer in ownership history."),
        ("initiateTransfer(...)","Citizen (Title Holder)","Puts in a request to transfer the vehicle. RTO must still approve for it to actually move."),
        ("updateVehicleState(...)","RTO Officer Only","Sets the vehicle to 'Scrapped'. Permanently locks the title. Cannot be undone."),
    ]
    for i, (fn, who, what) in enumerate(funcs):
        rw = tbl_f.rows[i+1]
        rw.cells[0].paragraphs[0].text = fn
        rw.cells[1].paragraphs[0].text = who
        rw.cells[2].paragraphs[0].text = what
        data_row(rw, i%2==1, fw)

    h2("7.3 Security Hook: No Privilege Owner Rule")
    body("We override OpenZeppelin's internal token transfer function (_update). Every time any token transfer "
         "is attempted — even directly using Ethereum's transferFrom — it first checks:")
    code_box("""if (to != address(0) && _isPrivilegedAddress(to)) {
    revert PrivilegedAddressCannotOwnVehicle(to);
}""")
    body("This means even if someone tries to bypass the web app and calls the contract directly, "
         "they cannot transfer a vehicle title to an RTO authority wallet. The rule is enforced at the code level.")

    h2("7.4 Web Application")
    body("The frontend is a React 18 + Vite application. It uses Ethers.js v6 to talk to the Ethereum network. "
         "Each user type has their own portal:")
    bullet("Anyone can enter a VIN or plate number and see if a vehicle is Active or Scrapped.", "Public Search Portal: ")
    bullet("Registered vehicle owners can view their digital title (shown as a realistic Indian HSRP plate), "
           "see the full ownership history, and request a transfer.", "Citizen Dashboard: ")
    bullet("Can register new vehicles and process approved transfer requests.", "RTO Registrar Portal: ")
    bullet("Can view vehicle details and mark vehicles as permanently scrapped.", "RTO Enforcement Officer Portal: ")
    doc.add_page_break()

    # ====================================================================
    # 8. PROJECT OUTPUTS & TEST RESULTS
    # ====================================================================
    h1("8. Project Outputs & Test Results")
    body("To prove the system works correctly and securely, we wrote 40 automated test cases using Hardhat and Chai. "
         "Every single test passed on the first full run. Here is a summary of what was tested:")

    h2("8.1 What Was Tested")
    tbl_t = doc.add_table(rows=9, cols=3)
    tbl_t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_border(tbl_t)
    ttw = [Inches(0.6), Inches(3.6), Inches(2.3)]
    for ci, hd in enumerate(["#","Test Group","What It Verifies"]):
        tbl_t.rows[0].cells[ci].paragraphs[0].text = hd
    hdr_row(tbl_t.rows[0], ttw)
    tgroups = [
        ("1","Role-Based Access Control & RTO Role Boundaries","Admin, Registrar, and Officer each have only what they need"),
        ("2","On-Chain Privacy (No Duplicate VIN / Plate)","Duplicate registrations are blocked at the hash level"),
        ("3","Two-Step Transfer Workflow","Citizens can request; only RTO can approve and execute"),
        ("4","Vehicle Scrappage & Terminal State","Scrapped vehicles cannot be transferred or re-registered"),
        ("5","Ownership History Tracking","All past owners are permanently recorded with timestamps"),
        ("6","Emergency Pause Controls","Admin can pause the system; no operations work while paused"),
        ("7","Event Privacy","Blockchain events contain only hashes — no raw personal data"),
        ("8","16 Negative Security Invariants","16 different attack scenarios are all blocked and properly rejected"),
    ]
    for i, (n, grp, what) in enumerate(tgroups):
        rw = tbl_t.rows[i+1]
        rw.cells[0].paragraphs[0].text = n
        rw.cells[1].paragraphs[0].text = grp
        rw.cells[2].paragraphs[0].text = what
        data_row(rw, i%2==1, ttw, [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    h2("8.2 Test Result Summary")
    code_box("""  Hardhat + Chai — Automated Test Suite
  ─────────────────────────────────────────────────
  Total Test Cases:   40
  Passed:             40
  Failed:              0
  Time Taken:          5 seconds
  ─────────────────────────────────────────────────
  All 16 negative security invariants: BLOCKED OK
  Ownership history tracking:          VERIFIED OK
  Scrappage terminal state:            ENFORCED OK
  Emergency pause:                     WORKING OK
  Role boundaries (Registrar/Officer): ENFORCED OK
  ─────────────────────────────────────────────────
  40 passing (5s)""")

    h2("8.3 Live Blockchain Test")
    body("We also ran a live demonstration script on a local Ethereum node (Chain ID: 31337). "
         "The script deployed the contract and ran through real transactions to prove everything works end-to-end:")
    code_box("""  Contract Deployed:  0x5FbDB2315678afecb367f032d93F642f64180aa3
  Admin Wallet:       0xf39Fd6e51aad88F6F4ce6aB8827279cffFb92266
  RTO Registrar:      0x70997970C51812dc3A010C7d01b50e0d17dc79C8
  RTO Officer:        0x15d34AAf54267DB7D7c367839AAf71A00a2C6A65
  ─────────────────────────────────────────────────────────────────
  [OK] Citizen wallet blocked from registering a vehicle
  [OK] Duplicate plate commitment rejected on-chain
  [OK] Ownership officially transferred via RTO Registrar
  [OK] RTO Officer blocked from calling transferOwnershipByRTO
  [OK] Citizen blocked from calling transferOwnershipByRTO
  [OK] Full ownership history verified on-chain (3 records)
  [OK] Vehicle scrapped by RTO Officer — state = SCRAPPED
  [OK] Transfer of scrapped vehicle blocked with error
  [OK] RTO wallet blocked from receiving a vehicle title
  [OK] Emergency pause activated — all operations blocked
  ─────────────────────────────────────────────────────────────────
  All proofs verified successfully on live chain.""")
    doc.add_page_break()

    # ====================================================================
    # 9. SKILLS DEVELOPED
    # ====================================================================
    h1("9. Skills Developed / Learning from this Project")
    body("Working on this project gave us practical experience with technologies and concepts that go well beyond "
         "what is covered in standard textbook exercises.")
    bullet("We learned to write real, deployable smart contracts using Solidity and OpenZeppelin. "
           "We understood how the Ethereum Virtual Machine (EVM) enforces rules automatically, without any middleman.",
           "Smart Contract Development: ")
    bullet("We practically applied Keccak-256 hashing to solve the real privacy problem of public blockchains. "
           "We now understand the difference between storing data publicly and proving it securely.",
           "Applied Cryptography: ")
    bullet("We designed and enforced role-based permissions at the code level. "
           "We understand why the Principle of Least Privilege is critical in any government or enterprise system.",
           "Access Control Design: ")
    bullet("We built a complete working web app that connects to a blockchain network in real time, "
           "handles wallet authentication, manages sessions, and reacts to live blockchain events.",
           "Full-Stack Web3 Development: ")
    bullet("We wrote 40 test cases before and alongside the code. We learned to actively think like an attacker "
           "to find and fix vulnerabilities before they become real problems.",
           "Test-Driven Security: ")

    # ====================================================================
    # 10. APPLICATIONS
    # ====================================================================
    h1("10. Applications of this Project")
    body("The architecture we built can be directly applied to solve real problems in India and worldwide:")
    bullet("State RTOs can use this as a backbone to eliminate forged RCs, duplicate registrations, "
           "and insider tampering across state borders.", "Government RTO Modernization: ")
    bullet("Platforms like Cars24, Spinny, and CarDekho can instantly verify a vehicle's complete ownership "
           "history and scrappage status without manual RTO paperwork.", "Used Vehicle Marketplaces: ")
    bullet("When an insurer declares a total loss, they can trigger scrappage on-chain. This prevents the "
           "salvage car from ever being resold to an unknowing buyer.", "Vehicle Insurance & Total-Loss Claims: ")
    bullet("Banks giving car loans can place a digital lien on the vehicle title during the loan period. "
           "The title cannot be transferred until the loan is fully repaid.", "Auto Finance & Hypothecation: ")

    # ====================================================================
    # 11. FUTURE IMPROVEMENTS
    # ====================================================================
    h1("11. Areas of Future Improvement")
    body("Our current system is a solid prototype. Here are areas where it can be extended:")
    bullet("A citizen could prove their vehicle has valid registration to a police officer "
           "without revealing their wallet address or identity, using zero-knowledge cryptography (zk-SNARKs).",
           "Zero-Knowledge Proof Integration: ")
    bullet("Vehicle inspection photos and certificates can be stored on IPFS (a distributed file network) "
           "instead of a central server, making the entire system fully decentralized.",
           "IPFS Storage for Documents: ")
    bullet("Vehicles with OBD-II chips and FASTag transponders can directly sign and submit "
           "tamper-proof odometer readings to the blockchain, ending odometer rollback fraud.",
           "IoT & FASTag Integration: ")
    bullet("Deploying on Ethereum Layer-2 networks like Polygon or Arbitrum would reduce transaction fees "
           "from a few dollars to a fraction of a rupee, making it practical for nationwide use.",
           "Layer-2 Deployment for Scale: ")

    # ====================================================================
    # 12. CONCLUSION
    # ====================================================================
    h1("12. Conclusion")
    body("This mini-project shows that blockchain technology can be used to solve real, everyday problems "
         "that affect ordinary citizens — not just for cryptocurrencies.")
    body("By building a vehicle registration system on Ethereum, we created records that nobody can secretly "
         "alter, a complete ownership history that anyone can verify, and a scrappage mechanism that "
         "permanently removes dangerous salvage cars from circulation — all while keeping citizens' "
         "personal details fully private through cryptographic hashing.")
    body("We also demonstrated that government administrative systems can be made tamper-proof by design, "
         "using smart contract rules instead of relying entirely on human honesty. The strict separation "
         "between the RTO Registrar and the RTO Enforcement Officer ensures that no single official has "
         "unchecked power to manipulate vehicle records.")
    body("Our 40 automated tests, all passing at 100%, and the live on-chain demonstration prove that "
         "this is not just a theoretical concept — it is a working, deployable system.")
    note_box("Blockchain does not just store data — it enforces rules. And rules enforced by code cannot "
             "be broken by any person, no matter how powerful they are. That is the real power of this technology.",
             "FINAL THOUGHT")

    # ── Save ────────────────────────────────────────────────────────────
    output = os.path.abspath(r"d:\BC project\VRS_Mini_Project_Report_Final.docx")
    doc.save(output)
    print(f"Saved: {output}  ({os.path.getsize(output)//1024} KB)")
    return output

if __name__ == "__main__":
    build()
